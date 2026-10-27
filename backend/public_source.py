"""Fonte mínima local, verificada antes de reutilizar o motor certificado.

Não lê Excel, não usa caches de resultados e não possui fallback de fonte.
O SHA é uma âncora de integridade da versão, não autorização de divulgação.
"""
from hashlib import sha256
import json
import math
from pathlib import Path
import re
from types import MappingProxyType

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter
from backend.motor import ActuarialEngine, BASE_RATES

DATA_PATH = Path(__file__).resolve().parent / 'data/atuarial_minimo.json'
MINIMUM_SHA256 = '85d3ff53f7b5675e5852efb7d5771aa22b0b6421a701f80eac5a584e3838954e'
ROOT_CELLS = ('BO5',) + tuple(f'D{r}' for r in range(6, 156)) + tuple(
    f'{get_column_letter(c)}157' for c in range(5, 63))
INPUT_CELLS = frozenset(f'C{r}' for r in range(6, 156))
REFERENCE = re.compile(r'\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?')


class SourceIntegrityError(ValueError):
    """Fonte ausente, incompleta ou divergente; mensagem sem detalhes internos."""


def dependency_closure(cells):
    """Inclui ambas as alternativas de IF; não depende de taxas amostradas."""
    pending = list(ROOT_CELLS)
    seen = set()
    while pending:
        address = pending.pop()
        if address in seen or address in INPUT_CELLS:
            continue
        seen.add(address)
        value = cells.get(address)
        if isinstance(value, str):
            if not value.startswith('='):
                raise SourceIntegrityError('Fonte atuarial inválida.')
            for token in Tokenizer(value).items:
                if token.type == 'OPERAND' and token.subtype == 'RANGE':
                    if not REFERENCE.fullmatch(token.value):
                        raise SourceIntegrityError('Fonte atuarial inválida.')
                    c1, r1, c2, r2 = range_boundaries(token.value)
                    if not (1 <= r1 <= r2 <= 157 and 1 <= c1 <= c2 <= 67):
                        raise SourceIntegrityError('Fonte atuarial inválida.')
                    pending.extend(f'{get_column_letter(c)}{r}'
                                   for r in range(r1, r2 + 1)
                                   for c in range(c1, c2 + 1))
    return seen


def unique_object(pairs):
    result = {}
    for key, value in pairs:
        if key in result:
            raise SourceIntegrityError('Fonte atuarial inválida.')
        result[key] = value
    return result


class PublicActuarialEngine(ActuarialEngine):
    """Altera apenas carregamento; evaluate/calculate são herdados sem cópia."""
    def __init__(self, source=DATA_PATH):
        try:
            self.path = Path(source).resolve()
            if self.path.stat().st_size > 1_000_000:
                raise SourceIntegrityError('Fonte atuarial inválida.')
            content = self.path.read_bytes()
            self.source_hash = sha256(content).hexdigest()
            if self.source_hash != MINIMUM_SHA256:
                raise SourceIntegrityError('Fonte atuarial indisponível ou inválida.')
            data = json.loads(content, object_pairs_hook=unique_object)
            if set(data) != {'schema', 'plans'} or type(data['schema']) is not int or data['schema'] != 1:
                raise SourceIntegrityError('Fonte atuarial inválida.')
            if set(data['plans']) != set(BASE_RATES):
                raise SourceIntegrityError('Fonte atuarial inválida.')
            plans = {}
            for plan, cells in data['plans'].items():
                for address, value in cells.items():
                    if not re.fullmatch(r'[A-Z]{1,3}\d+', address):
                        raise SourceIntegrityError('Fonte atuarial inválida.')
                    if value is not None and not (
                        type(value) in (int, float) and math.isfinite(value)
                        or isinstance(value, str) and value.startswith('=')):
                        raise SourceIntegrityError('Fonte atuarial inválida.')
                if dependency_closure(cells) != set(cells):
                    raise SourceIntegrityError('Fonte atuarial incompleta ou excedente.')
                plans[plan] = MappingProxyType(cells)
            self._plans = MappingProxyType(plans)
        except (OSError, ValueError, TypeError, KeyError, AttributeError) as exc:
            raise SourceIntegrityError('Fonte atuarial indisponível ou inválida.') from None
