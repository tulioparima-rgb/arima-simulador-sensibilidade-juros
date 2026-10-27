"""Avaliação restrita das fórmulas do fluxo aprovado, sem utilizar caches.

Taxas são percentuais anuais (4.91 significa 4,91%, não 491%).
O avaliador não executa código proveniente da planilha: AST e funções permitidas
são explicitamente enumeradas e qualquer referência não suportada falha fechada.
"""
from __future__ import annotations

import ast
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from hashlib import sha256
from io import BytesIO
from pathlib import Path
import math
import re

from openpyxl import load_workbook
from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import range_boundaries, get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WORK_COPY = ROOT / 'outputs/investigacao_vinculos/Fluxo_investigacao.xlsx'
SOURCE_SHA256 = '692efc67665d54d0c570d880e476597f7c05ae9fea13a49bce0c235e91863b3c'
BASE_RATES = {'Previdenciário': Decimal('4.91'), 'Financeiro': Decimal('5.42')}


class InvalidRate(ValueError):
    """Entrada não satisfaz o contrato de taxa."""


class FormulaError(ValueError):
    """Fórmula/referência não suportada; nunca substituir silenciosamente por zero."""


def validate_rate(value: object) -> Decimal:
    if isinstance(value, bool) or not isinstance(value, (str, int, float, Decimal)):
        raise InvalidRate('Informe taxa numérica entre 0,00 e 20,00, com até duas casas.')
    if isinstance(value, str):
        value = value.strip()
        if not re.fullmatch(r'[+]?\d+(?:[.,]\d{1,2})?', value):
            raise InvalidRate('Formato inválido: use percentual sem símbolo %, até duas casas.')
        value = value.replace(',', '.')
    try:
        rate = Decimal(str(value))
    except (InvalidOperation, ValueError) as exc:
        raise InvalidRate('Taxa não numérica.') from exc
    if not rate.is_finite() or not Decimal(0) <= rate <= Decimal(20):
        raise InvalidRate('Taxa fora do intervalo de 0,00 a 20,00.')
    if rate != rate.quantize(Decimal('0.01')):
        raise InvalidRate('A taxa deve respeitar incrementos de 0,01 ponto percentual.')
    return rate.quantize(Decimal('0.01'))


def excel_round(value: float, digits: int) -> float:
    """ROUND decimal, empate afastado de zero; não usa round bancário do Python.

    Operações intermediárias usam binary64 como Excel. A fronteira ROUND
    converte a representação decimal de 15 dígitos significativos do resultado,
    preservando os pontos de arredondamento do arquivo. Essa fronteira evita
    que resíduos binary64 abaixo do empate alterem o ROUND observado no Excel.
    Compatibilidade verificada nesta fonte/domínio; não é ROUND Excel universal.
    """
    if not math.isfinite(value):
        raise FormulaError('Valor não finito no arredondamento.')
    with localcontext() as context:
        context.prec = 40
        decimal_value = Decimal(format(value, '.15g'))
        return float(decimal_value.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP))


def money(value: float) -> Decimal:
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)


@dataclass(frozen=True)
class Result:
    plan: str
    rate: Decimal
    guaranteed_assets: Decimal
    gross_assets: Decimal
    exclusion: Decimal
    adjusted_assets: Decimal
    liabilities: Decimal
    result: Decimal


class _Evaluation:
    def __init__(self, cells: dict, rate: Decimal):
        self.cells = cells
        self.rate = float(rate)
        self.values: dict[str, float] = {}
        self.active: set[str] = set()

    def cell(self, reference: str):
        reference = reference.replace('$', '')
        if not re.fullmatch(r'[A-Z]{1,3}\d+', reference):
            raise FormulaError(f'Referência não suportada: {reference}')
        if reference in self.values:
            return self.values[reference]
        if reference in self.active:
            raise FormulaError(f'Ciclo detectado: {reference}')
        if reference.startswith('C') and reference[1:].isdigit() and 6 <= int(reference[1:]) <= 155:
            return self.rate
        raw = self.cells.get(reference)
        if raw is None:
            return 0.0  # Excel trata referência numérica a célula vazia como zero.
        if isinstance(raw, (int, float)):
            return float(raw)
        if not isinstance(raw, str) or not raw.startswith('='):
            raise FormulaError(f'Célula não numérica em cálculo: {reference}')
        self.active.add(reference)
        try:
            tokens = []
            for token in Tokenizer(raw).items:
                if token.type == 'OPERAND' and token.subtype == 'RANGE':
                    tokens.append(f'REF({token.value!r})')
                elif token.type == 'OPERATOR-INFIX':
                    tokens.append({'^':'**', '=':'==', '<>':'!='}.get(token.value,token.value))
                else:
                    tokens.append(token.value)
            expression = ast.parse(''.join(tokens), mode='eval').body
            result = self.node(expression)
            if not isinstance(result, (int,float)) or not math.isfinite(result):
                raise FormulaError(f'Resultado inválido em {reference}')
            self.values[reference] = result
            return result
        except (SyntaxError, ZeroDivisionError, OverflowError) as exc:
            raise FormulaError(f'Erro em {reference}: {raw}') from exc
        finally:
            self.active.remove(reference)

    def reference(self, value):
        if not isinstance(value,str) or not re.fullmatch(r'\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?',value):
            raise FormulaError(f'Referência externa ou dinâmica rejeitada: {value}')
        if ':' not in value:
            return self.cell(value)
        c1,r1,c2,r2 = range_boundaries(value)
        return [self.cell(f'{get_column_letter(c)}{r}') for r in range(r1,r2+1) for c in range(c1,c2+1)]

    def node(self, n):
        if isinstance(n, ast.Constant) and isinstance(n.value,(int,float,str)):
            return n.value
        if isinstance(n, ast.UnaryOp):
            value=self.node(n.operand)
            if isinstance(n.op,ast.USub): return -value
            if isinstance(n.op,ast.UAdd): return value
        if isinstance(n,ast.BinOp):
            left,right=self.node(n.left),self.node(n.right)
            if isinstance(n.op,ast.Add): return left+right
            if isinstance(n.op,ast.Sub): return left-right
            if isinstance(n.op,ast.Mult): return left*right
            if isinstance(n.op,ast.Div): return left/right
            if isinstance(n.op,ast.Pow): return left**right
        if isinstance(n,ast.Compare) and len(n.ops)==1:
            left,right=self.node(n.left),self.node(n.comparators[0]); op=n.ops[0]
            if isinstance(op,ast.Eq): return left==right
            if isinstance(op,ast.NotEq): return left!=right
            if isinstance(op,ast.Lt): return left<right
            if isinstance(op,ast.LtE): return left<=right
            if isinstance(op,ast.Gt): return left>right
            if isinstance(op,ast.GtE): return left>=right
        if isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and not n.keywords:
            name=n.func.id
            if name=='IF' and len(n.args)==3:
                return self.node(n.args[1] if self.node(n.args[0]) else n.args[2])
            args=[self.node(a) for a in n.args]
            if name=='REF' and len(args)==1: return self.reference(args[0])
            if name=='ROUND' and len(args)==2: return excel_round(args[0],int(args[1]))
            if name=='SUM': return sum(x for a in args for x in (a if isinstance(a,list) else [a]))
            if name=='AND': return all(args)
            if name=='SUMPRODUCT' and len(args)==2:
                if not all(isinstance(a,list) for a in args) or len(args[0])!=len(args[1]):
                    raise FormulaError('SUMPRODUCT com intervalos incompatíveis.')
                return sum(a*b for a,b in zip(*args))
        raise FormulaError(f'Operação não suportada: {ast.dump(n)}')


class ActuarialEngine:
    """Fonte única verificada; cada avaliação tem estado isolado por plano/taxa."""
    def __init__(self, workbook: Path | str = WORK_COPY):
        self.path = Path(workbook).resolve()
        if self.path == (ROOT/'Fluxo.xlsx').resolve():
            raise ValueError('Use cópia isolada, nunca o original para execução.')
        content = self.path.read_bytes()
        self.source_hash = sha256(content).hexdigest()
        if self.source_hash != SOURCE_SHA256:
            raise ValueError('Fonte diverge do SHA-256 homologado para esta versão.')
        book=load_workbook(BytesIO(content),data_only=False,keep_links=False)
        self._plans={s:{c.coordinate:c.value for row in book[s] for c in row if c.value is not None} for s in BASE_RATES}
        book.close()

    def evaluate(self, plan: str, rate: object) -> tuple[Result, dict[str,float]]:
        if plan not in self._plans:
            raise ValueError('Plano desconhecido.')
        valid_rate=validate_rate(rate)
        evaluation=_Evaluation(self._plans[plan],valid_rate)
        # Pré-avaliar sequência recursiva para não depender da profundidade da pilha.
        for row in range(6,156):
            evaluation.cell(f'D{row}')
        # Todas as rubricas e subtotais de VP, na ordem espacial do arquivo.
        for col in range(5,63):
            address=f'{get_column_letter(col)}157'
            if address in self._plans[plan]: evaluation.cell(address)
        assets,gross,exclusion,liabilities=(money(evaluation.cell(c)) for c in ('BO5','AR157','AQ157','BJ157'))
        adjusted=gross-exclusion if plan=='Financeiro' else gross
        result=Result(plan,valid_rate,assets,gross,exclusion,adjusted,liabilities,assets+adjusted-liabilities)
        return result,dict(evaluation.values)

    def calculate(self, plan: str, rate: object) -> Result:
        return self.evaluate(plan,rate)[0]

    def calculate_both(self, previdenciario: object='4.91', financeiro: object='5.42') -> dict[str,Result]:
        return {'Previdenciário':self.calculate('Previdenciário',previdenciario),'Financeiro':self.calculate('Financeiro',financeiro)}
